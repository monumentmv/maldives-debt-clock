"""
Ministry of Finance lists of active external loans and active sovereign guarantees.

Every new version of each file is kept in inputs/mof/, named by its "as of" date, so the history builds up
over time (the Ministry replaces the files on its website). Every version is read into
data/lenders/raw.json, with the IMF exchange rates on each version's "as of" date in data/lenders/fx.json
(the IMF data used is kept in inputs/imf/).
build_lenders.py then turns that into lenders.json for the home page.

Runs every morning in GitHub Actions. To run it yourself:
    pip install requests openpyxl
    python fetch_lenders.py
"""
import io
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import requests

HERE = Path(__file__).resolve().parent
INPUTS = HERE / "inputs" / "mof"
OUT = HERE / "data" / "lenders"
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36"}
# the files' current addresses on finance.gov.mv (the Ministry replaces the file behind each address)
FILES = {
    "guarantees": "https://www.finance.gov.mv/public/attachments/TNxIzceEndB8ZaPZbzjn9aMXG5wY4HFgJG473Rwy.xlsx",
    "loans": "https://www.finance.gov.mv/public/attachments/uhclBk3Vdbveshqh4b9u5jsUIqEfSOBch5ewr32p.xlsx",
}
MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
PEGS = {"USD": 1.0, "SAR": 3.75, "AED": 3.6725}   # currencies fixed to the US dollar


def as_of(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=6, values_only=True):
            for c in row:
                m = re.search(r"as of\s+(\d{1,2})\s+([A-Za-z]{3})[a-z]*\s+(\d{4})", str(c or ""), re.I)
                if m:
                    return date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1))).isoformat()
    return None


def day(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    m = re.match(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", str(v or ""))
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat() if m else None


def rows_of(ws):
    """The table under the header row that starts with '#', as dicts keyed by simplified header names."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and str(r[0] or "").strip() == "#" or (len(r) > 1 and str(r[0] or "").strip() == "" and str(r[1] or "").strip() == "#"):
            hdr = [re.sub(r"[^a-z%#]+", " ", str(h or "").lower()).strip() for h in r]
            out = []
            for r2 in rows[i + 1:]:
                if not any(c not in (None, "") for c in r2):
                    break
                d = {hdr[j]: r2[j] for j in range(len(hdr)) if hdr[j]}
                if not isinstance(d.get("#"), (int, float)):
                    break
                out.append(d)
            return out
    return []


def pick(d, *names):
    for n in names:
        for k, v in d.items():
            if k.startswith(n):
                return v
    return None


def parse(kind, content):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    asof = as_of(wb)
    items = []
    for ws in wb.worksheets:
        group = "private" if "private" in ws.title.lower() else ("soe" if kind == "guarantees" else "loan")
        for d in rows_of(ws):
            amt = pick(d, "guaranteed amount", "amount")
            items.append({
                "group": group, "date": day(pick(d, "loan date", "date")),
                "creditor": str(pick(d, "creditor") or "").strip().replace("Male'", "Malé").rstrip("'").strip(),
                "beneficiary": str(pick(d, "beneficiary") or "").strip() or None,
                "purpose": re.sub(r"\s+", " ", str(pick(d, "purpose", "project") or "")).strip(),
                "currency": str(pick(d, "curren") or "").strip().upper(),
                "amount": float(amt) if isinstance(amt, (int, float)) else None,
                "status": str(pick(d, "project implementation", "implementation status", "implementation") or "").strip() or None,
            })
    return asof, items


# IMF exchange rates (dataset ER on the IMF's data service): end-of-month rates, national currency per US dollar
IMF = "https://api.imf.org/external/sdmx/2.1/data/IMF.STA,ER/{key}?startPeriod={start}"
IMF_SAVE = HERE / "inputs" / "imf"
IMF_COUNTRY = {"CNY": "CHN", "INR": "IND", "JPY": "JPN", "KWD": "KWT", "SAR": "SAU", "AED": "ARE", "EUR": "DEU",
               "GBP": "GBR", "CHF": "CHE", "KRW": "KOR", "SGD": "SGP", "AUD": "AUS", "XDR": "USA"}


def month_end(d):
    """The last month-end on or before day d, as (YYYY-MM, YYYY-MM-DD)."""
    x = date.fromisoformat(d)
    if (x + timedelta(days=1)).month == x.month:   # not the last day of its month: use the month before
        x = x.replace(day=1) - timedelta(days=1)
    return f"{x:%Y-%m}", x.isoformat()


def imf_series(key, start):
    """[(country, indicator, indicator name, {YYYY-MM: value})] for an SDMX key on the IMF's data service."""
    r = requests.get(IMF.format(key=key, start=start), headers={**H, "Accept": "application/json"}, timeout=90)
    r.raise_for_status()
    IMF_SAVE.mkdir(parents=True, exist_ok=True)
    (IMF_SAVE / f"ER_{key.split('.')[0][:40].replace('+', '_')}_{start}.json").write_text(r.text)   # kept, so every rate used can be checked
    j = r.json()
    dims = j["structure"]["dimensions"]["series"]
    times = [t["id"].replace("-M", "-") for t in j["structure"]["dimensions"]["observation"][0]["values"]]
    pos = {dm["id"]: i for i, dm in enumerate(dims)}
    out = []
    for k, ser in j["dataSets"][0]["series"].items():
        idx = [int(i) for i in k.split(":")]
        val = lambda dim: dims[pos[dim]]["values"][idx[pos[dim]]]
        obs = {times[int(t)]: float(o[0]) for t, o in ser["observations"].items() if o and o[0] not in (None, "")}
        out.append((val("COUNTRY")["id"], val("INDICATOR")["id"], val("INDICATOR")["name"], obs))
    return out


def imf_rates(d, need):
    """Units of each currency per US dollar at the last month-end on or before day d, from the IMF."""
    month, day_ = month_end(d)
    rates = {}
    plain = [c for c in need if c in IMF_COUNTRY and c != "XDR"]
    try:
        if plain:
            back = {v: k for k, v in IMF_COUNTRY.items() if k in plain}
            for country, ind, name, obs in imf_series("+".join(IMF_COUNTRY[c] for c in plain) + ".XDC_USD.EOP_RT.M", month):
                if month in obs and country in back:
                    rates[back[country]] = obs[month]
        if "XDR" in need:   # the SDR: from the US dollar's rate against it
            for country, ind, name, obs in imf_series("USA..EOP_RT.M", month):
                print(f"  IMF USA indicator {ind}: {name} = {obs.get(month)}")
                n = name.lower()
                if month in obs and "sdr" in n and "per" in n and "eur" not in n:
                    v = obs[month]
                    rates["XDR"] = 1 / v if n.index("per") < n.index("sdr") else v   # "... per SDR" is US$ per SDR
                    break
    except Exception as e:
        print(f"  IMF exchange rates for {month} not available ({e})")
    return rates, day_


def fx_on(d, currencies, usd_mvr, before=None):
    """Units of each currency per US dollar for day d: the IMF's end-of-month rate. MVR uses the MMA rate."""
    need = sorted(c for c in currencies if c not in ("USD", "MVR"))
    rates, notes = {"USD": 1.0} if "USD" in currencies else {}, {}
    imf, used = imf_rates(d, need) if need else ({}, d)
    for c in need:
        v, old = imf.get(c), (before or {}).get(c)
        if v and old and abs(v / old - 1) > .2:
            print(f"  check {c}: IMF {v} against {old} before")
        if v:
            rates[c] = v
    missing = [c for c in need if c not in rates]
    for c in missing:   # not in the IMF's data yet: fixed pegs, else the open currency-api dataset
        if c in PEGS:
            rates[c], notes[c] = PEGS[c], f"fixed to the US dollar at {PEGS[c]}"
            continue
        for url in (f"https://cdn.jsdelivr.net/npm/@fawazahmed0/currency-api@{d}/v1/currencies/usd.json",
                    f"https://{d}.currency-api.pages.dev/v1/currencies/usd.json"):
            try:
                rates[c] = float(requests.get(url, timeout=60).json()["usd"][c.lower()])
                notes[c] = "not yet in the IMF's data: open currency-api dataset"
                break
            except Exception:
                pass
    if "MVR" in currencies and usd_mvr:
        rates["MVR"], notes["MVR"] = usd_mvr, "MMA reference rate"
    from_imf = bool(imf)
    when = f"{date.fromisoformat(used):%-d %B %Y}"
    return {"date": d, "rate_date": used, "per_usd": rates, "notes": notes, "imf": from_imf,
            "source": f"International Monetary Fund (IMF) exchange rates, end of month ({when})" if from_imf else None}


def mma_rate(d):
    try:
        pts = json.load(open(HERE / "data.json"))["series"]["usd_rate"]["points"]
        before = [p for p in pts if p["date"] <= d] or pts
        return before[-1]["value"]
    except Exception:
        return None


def main():
    INPUTS.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    for kind, url in FILES.items():
        try:
            r = requests.get(url, headers=H, timeout=90)
            r.raise_for_status()
            if not r.content.startswith(b"PK"):
                raise RuntimeError("not an Excel file")
            asof, items = parse(kind, r.content)
            if not asof or not items:
                raise RuntimeError("couldn't find the 'As of' date or the table")
        except Exception as e:
            print(f"{kind}: couldn't use the Ministry's file ({e}). Keeping the saved versions.")
            continue
        path = INPUTS / f"{kind}_{asof}.xlsx"
        if path.exists() and path.read_bytes() == r.content:
            print(f"{kind}: as of {asof}, no change.")
        else:
            path.write_bytes(r.content)
            print(f"{kind}: as of {asof}, {len(items)} rows, saved as {path.relative_to(HERE)}.")

    raw, fx = {"loans": {}, "guarantees": {}}, {}
    try:
        fx = json.load(open(OUT / "fx.json"))
    except Exception:
        pass
    for f in sorted(INPUTS.glob("*.xlsx")):
        kind = f.name.split("_")[0]
        asof, items = parse(kind, f.read_bytes())
        raw[kind][asof] = items
    # exchange rates on each "as of" day, for every currency used that day (looked up once, then kept)
    for asof in sorted({a for k in raw for a in raw[k]}):
        cur = {i["currency"] for k in raw for i in raw[k].get(asof, []) if i["currency"]}
        if not fx.get(asof, {}).get("imf") or not cur <= set(fx[asof].get("per_usd", {})):
            fx[asof] = fx_on(asof, cur, mma_rate(asof), fx.get(asof, {}).get("per_usd"))
            print(f"  exchange rates for {asof}: " + ", ".join(f"{c} {v:.4f}" for c, v in sorted(fx[asof]['per_usd'].items())))
            missing = sorted(cur - set(fx[asof]["per_usd"]))
            if missing:
                print(f"  no exchange rate for {', '.join(missing)} on {asof}")
    json.dump(raw, open(OUT / "raw.json", "w"), ensure_ascii=False, indent=0)
    json.dump(fx, open(OUT / "fx.json", "w"), indent=1)
    creditors = sorted({i["creditor"] for k in raw for v in raw[k].values() for i in v})
    print(f"Versions saved: {', '.join(f'{k} {sorted(v)}' for k, v in raw.items())}")
    print(f"{len(creditors)} creditors:\n  " + "\n  ".join(creditors))


if __name__ == "__main__":
    main()
