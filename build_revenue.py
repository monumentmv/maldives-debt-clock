"""
Turns MIRA's monthly revenue spreadsheets into revenue.json for the revenue page.

Reads every .xlsx file in inputs/mira/ and recognises the MIRA files by their title:
  "REVENUE SERIES (MONTHLY) - TOTAL REVENUE COLLECTION"  all collections, in rufiyaa
  "REVENUE SERIES (MONTHLY) - USD REVENUE COLLECTION"    the part paid in US dollars
  "REVENUE SERIES (WEEKLY) - REVENUE COLLECTION"          provisional weekly collections, in rufiyaa and in US dollars (optional)

Runs in GitHub Actions whenever a file in inputs/mira/ changes. To run it yourself:
    pip install openpyxl
    python build_revenue.py
"""
import csv
import json
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Please run:  pip install openpyxl")

HERE = Path(__file__).resolve().parent
INPUTS = HERE / "inputs" / "mira"
OUT = HERE / "revenue.json"
MON = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]

GROUPS = [
    ("tourism_gst", "Tourism GST"),
    ("gst_other", "GST outside tourism"),
    ("green", "Green tax"),
    ("airport", "Airport taxes and fees"),
    ("resort_land", "Resort land and leases"),
    ("business", "Business profit taxes"),
    ("personal", "Personal income taxes"),
    ("tourism_tax", "Tourism tax (ended 2014)"),
    ("other", "Other taxes and fees"),
]
GROUP_OF = {
    "gst (tourism sector)": "tourism_gst",
    "gst (non-tourism sector)": "gst_other",
    "green tax": "green",
    "airport service charge": "airport",
    "airport development fee": "airport",
    "departure tax": "airport",
    "tourism land rent": "resort_land",
    "lease period extension fee": "resort_land",
    "land acquisition and conversion fee": "resort_land",
    "construction period extension fee": "resort_land",
    "corporate income tax": "business",
    "business profit tax": "business",
    "bank profit tax": "business",
    "withholding tax": "business",
    "non-resident withholding tax": "business",
    "personal income tax": "personal",
    "employee withholding tax": "personal",
    "tourism tax": "tourism_tax",
}
# lines shown on their own inside a group
DETAIL = {"tourism land rent": "Tourism land rent", "lease period extension fee": "Lease period extensions",
          "land acquisition and conversion fee": "Land acquisition and conversion",
          "construction period extension fee": "Construction period extensions"}


def norm(name):
    n = re.sub(r"\s+", " ", str(name)).strip().lower()
    return n.replace("goods and services tax", "gst")


def read_sheet(path):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    title = str(rows[0][0] or "").upper()
    kind = "usd" if "USD REVENUE" in title else "all" if "TOTAL REVENUE" in title else None
    if not kind:
        return None
    years, months = rows[1], rows[2]
    cols, y = {}, None
    for j in range(1, len(months)):
        if years[j] not in (None, ""):
            y = int(years[j])
        m = str(months[j] or "").strip().lower()[:3]
        if y and m in MON:
            cols[j] = f"{y}-{MON.index(m) + 1:02d}"
    items, total, note = {}, None, ""
    for r in rows[3:]:
        label = r[0]
        if label is None:
            continue
        if str(label).lower().startswith("updated on"):
            note = str(label).strip()
            continue
        vals = {cols[j]: (float(r[j]) if isinstance(r[j], (int, float)) else None) for j in cols}
        if total is not None:
            continue  # zakat and donations sit below TOTAL and are not part of it
        if norm(label) == "total":
            total = vals
            continue
        items[str(label).strip()] = vals
    if total is None:
        sys.exit(f"{path.name}: no TOTAL row found")
    return {"kind": kind, "items": items, "total": total, "note": note, "file": path.name}


# the weekly file's lines, sorted into the same groups as the monthly figures
WEEK_GROUP = {"corporate income tax/business profit tax": "business", "non-resident withholding tax": "business",
              "gst (tourism sector)": "tourism_gst", "gst (non-tourism sector)": "gst_other", "green tax": "green",
              "departure tax / airport service charge": "airport", "airport development fee": "airport",
              "tourism land rent": "resort_land", "other taxes and fees": "other"}
MONTH_NO = {m: i + 1 for i, m in enumerate(MON)}


def week_dates(head):
    """'26th July 2026 to 1st August 2026', '9th Aug to 15th Aug 2020', '1st July - 04th July 2020' -> (first day, last day)"""
    parts = re.findall(r"(\d{1,2})(?:st|nd|rd|th)?(?:\s+([A-Za-z]{3})[a-z]*\.?)?(?:\s+(\d{4}))?", str(head))
    if len(parts) != 2:
        return None
    (d1, m1, y1), (d2, m2, y2) = parts
    m1 = m1 or m2                       # '18th to 24th Apr 2021
    if not y2 or m1.lower() not in MONTH_NO or m2.lower() not in MONTH_NO:
        return None
    end = date(int(y2), MONTH_NO[m2.lower()], int(d2))
    start = date(int(y1 or y2), MONTH_NO[m1.lower()], int(d1))
    if start > end:                     # a week running over new year, with the year written only once
        start = start.replace(year=start.year - 1)
    if (end - start).days > 6:          # a slip in the heading (MIRA wrote "04th July to 10th August 2024"): a week is the 7 days to its last day
        start = end - timedelta(days=6)
    return start, end


def read_weekly(path):
    """MIRA's weekly file: one sheet in rufiyaa, one for the part paid in US dollars. None if it isn't that file."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows or "WEEKLY" not in str(rows[0][0] or "").upper():
            return None
        kind = "usd" if "USD" in ws.title.upper() else "mvr"
        weeks = {}
        for j, h in enumerate(rows[1][1:], start=1):
            w = week_dates(h) if h else None
            if w:
                weeks[j] = w
            elif h:
                print(f"{path.name} ({ws.title}): couldn't read the dates in column heading '{h}', skipping that week.")
        lines, total = {}, None
        for r in rows[2:]:
            label = str(r[0] or "").strip()
            if not label or label.lower().startswith("provisional"):
                continue
            vals = {weeks[j][1].isoformat(): (float(r[j]) if isinstance(r[j], (int, float)) else 0.0) for j in weeks if j < len(r)}
            if label.upper() == "TOTAL":
                total = vals
            else:
                lines[label] = vals
        if total is None:
            sys.exit(f"{path.name} ({ws.title}): no TOTAL row found")
        out[kind] = {"weeks": {w[1].isoformat(): w[0].isoformat() for w in weeks.values()}, "lines": lines, "total": total}
    return out if "mvr" in out else None


def weekly_json(W, name):
    """the weekly figures for revenue.json, grouped like the monthly ones, with checks that the lines add up"""
    ends = sorted(W["mvr"]["weeks"])
    problems, unknown = [], set()
    def grouped(S):
        g = {}
        for label, vals in S["lines"].items():
            k = WEEK_GROUP.get(re.sub(r"\s+", " ", label).lower())
            if not k:
                unknown.add(label); k = "other"
            arr = g.setdefault(k, [0.0] * len(ends))
            for i, e in enumerate(ends):
                arr[i] += vals.get(e, 0.0)
        return {k: [round(v) for v in arr] for k, arr in g.items()}
    for kind, S in W.items():
        for e in S["weeks"]:
            s, t = sum(v.get(e, 0.0) for v in S["lines"].values()), S["total"].get(e, 0.0)
            if abs(s - t) > max(1.0, abs(t) * 1e-6):
                problems.append(f"weekly {kind}, week to {e}: lines add up to {s:,.0f} but TOTAL says {t:,.0f}")
    U = W.get("usd")
    usd_ends = sorted(U["weeks"]) if U else []
    data = {"file": name, "start": [W["mvr"]["weeks"][e] for e in ends], "end": ends,
            "mvr": grouped(W["mvr"]), "total_mvr": [round(W["mvr"]["total"].get(e, 0.0)) for e in ends],
            "usd_from": usd_ends[0] if usd_ends else None}
    if U:
        g = grouped({"lines": {l: v for l, v in U["lines"].items()}})
        data["usd"] = {k: [None if e < usd_ends[0] else v for e, v in zip(ends, arr)] for k, arr in g.items()}
        data["total_usd"] = [None if e < usd_ends[0] else round(U["total"].get(e, 0.0)) for e in ends]
        missing = [e for e in usd_ends if e not in W["mvr"]["weeks"]]
        if missing:
            problems.append(f"weekly usd: {len(missing)} weeks have no rufiyaa column ({missing[0]} first)")
    gaps = [b for a, b in zip(ends, ends[1:]) if (date.fromisoformat(b) - date.fromisoformat(a)).days != 7]
    if gaps:
        problems.append(f"weekly: {len(gaps)} weeks don't follow on 7 days after the one before ({gaps[0]} first)")
    if unknown:
        print("Weekly lines counted as other taxes and fees: " + ", ".join(sorted(unknown)))
    return data, problems


# MIRA's weekly lines next to the same taxes in Ministry of Finance's weekly fiscal report (data/wfd/wfd_long.csv)
VS_FINANCE = [
    ("Tourism GST", ["gst (tourism sector)"], ["tourism goods and services tax"]),
    ("GST outside tourism", ["gst (non-tourism sector)"], ["general goods and services tax"]),
    ("Green tax", ["green tax"], ["green tax"]),
    ("Business profit tax", ["corporate income tax/business profit tax"], ["corporate income tax", "business profit tax"]),
    ("Non-resident withholding tax", ["non-resident withholding tax"], ["non-resident withholding tax"]),
    ("Resort land rent", ["tourism land rent"], ["rent from resorts"]),
    ("Airport taxes and fees", ["departure tax / airport service charge", "airport development fee"], ["airport service charges / departure tax", "airport development fee"]),
]
WFD_LONG = HERE / "data" / "wfd" / "wfd_long.csv"


def vs_finance(W, year, upto):
    """MIRA's weekly collections for `year` up to the week ending `upto`, next to Finance's report closest before that day.
    A week that straddles new year counts only for its days in the year."""
    if not WFD_LONG.exists():
        return None
    fin = {}
    with open(WFD_LONG, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["table"] == "revenue" and r["as_at"].startswith(str(year)) and r["as_at"] <= upto:
                fin.setdefault(r["as_at"], {})[re.sub(r"\s+", " ", r["label"]).strip().lower()] = r["this_year"]
    if not fin:
        return None
    at = max(fin)
    if (date.fromisoformat(upto) - date.fromisoformat(at)).days > 7:
        return None
    S, jan1, cap = W["mvr"], date(year, 1, 1), min(date(year, 12, 31), date.fromisoformat(upto))
    lines = {re.sub(r"\s+", " ", k).strip().lower(): v for k, v in S["lines"].items()}
    def mira(names):
        t = 0.0
        for e, st in S["weeks"].items():
            e_, s_ = date.fromisoformat(e), date.fromisoformat(st)
            inside = (min(e_, cap) - max(s_, jan1)).days + 1      # days of this week that fall in the period
            if inside <= 0:
                continue
            share = inside / ((e_ - s_).days + 1)
            t += share * sum(lines.get(n, {}).get(e, 0.0) for n in names)
        return round(t / 1e6, 1)
    def finance(names):
        vals = [fin[at].get(n) for n in names]
        return round(sum(float(v) for v in vals if v not in (None, "")), 1) if any(v not in (None, "") for v in vals) else None
    rows = [{"label": l, "mira": mira(m), "finance": finance(f)} for l, m, f in VS_FINANCE]
    rows = [r for r in rows if r["finance"]]
    return {"year": year, "mira_to": upto, "finance_at": at, "rows": rows}


def main():
    files = sorted(INPUTS.glob("*.xlsx"))
    if not files:
        print("No MIRA files in inputs/mira/, nothing to do.")
        return
    found, weekly = {}, None
    for f in files:
        w = read_weekly(f)
        if w:
            if weekly:
                sys.exit(f"Two weekly files in inputs/mira/ ({weekly[1]}, {f.name}). Delete the older one.")
            weekly = (w, f.name)
            continue
        s = read_sheet(f)
        if s:
            if s["kind"] in found:
                sys.exit(f"Two {s['kind']} files in inputs/mira/ ({found[s['kind']]['file']}, {f.name}). Delete the older one.")
            found[s["kind"]] = s
        else:
            print(f"Skipping {f.name}, not a recognised MIRA monthly revenue file.")
    if "all" not in found:
        sys.exit("The total revenue collection file is missing from inputs/mira/.")
    A, U = found["all"], found.get("usd")

    # months run from the first to the last month with a total above zero
    months = sorted(m for m, v in A["total"].items() if v)
    first, last = months[0], months[-1]
    months = sorted(m for m in A["total"] if first <= m <= last)
    usd_months = sorted(m for m, v in (U["total"] if U else {}).items() if v and m <= last)
    usd_from = usd_months[0] if usd_months else None

    problems = []
    for name, S in (("all", A), ("usd", U)):
        if not S:
            continue
        for m in months:
            t = S["total"].get(m)
            if t is None or (name == "usd" and (not usd_from or m < usd_from)):
                continue
            s = sum(v[m] or 0 for v in S["items"].values())
            if abs(s - t) > max(1.0, abs(t) * 1e-6):
                problems.append(f"{name} file, {m}: lines add up to {s:,.0f} but TOTAL says {t:,.0f}")
    if problems:
        print("Warning, totals don't match:\n  " + "\n  ".join(problems[:20]))

    def by_group(S, start):
        out = {k: [0.0] * len(months) for k, _ in GROUPS}
        detail = {label: [0.0] * len(months) for label in DETAIL.values()}
        unknown = set()
        for name, vals in S["items"].items():
            n = norm(name)
            g = GROUP_OF.get(n, "other")
            if g == "other" and n not in ("others",):
                unknown.add(name)
            for i, m in enumerate(months):
                v = vals.get(m) or 0.0
                out[g][i] += v
                if n in DETAIL:
                    detail[DETAIL[n]][i] += v
        for arr in list(out.values()) + list(detail.values()):
            for i, m in enumerate(months):
                arr[i] = None if (start and m < start) else round(arr[i], 2)
        return out, detail, unknown

    mvr, mvr_detail, other_all = by_group(A, None)
    usd, usd_detail, other_usd = (by_group(U, usd_from) if U else (None, None, set()))
    total_mvr = [round(A["total"].get(m) or 0.0, 2) for m in months]
    total_usd = [None if (not U or not usd_from or m < usd_from) else round(U["total"].get(m) or 0.0, 2) for m in months]

    data = {
        "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": "Maldives Inland Revenue Authority, monthly revenue collection",
        "source_note": A["note"],
        "files": [A["file"]] + ([U["file"]] if U else []),
        "start": months[0],
        "end": months[-1],
        "usd_start": usd_from,
        "months": months,
        "groups": [{"key": k, "label": l} for k, l in GROUPS],
        "total": {"mvr": total_mvr, "usd": total_usd},
        "mvr": mvr,
        "usd": usd,
        "detail": {"mvr": mvr_detail, "usd": usd_detail},
        "other_lines": {"all": sorted(other_all), "usd": sorted(other_usd)},
        "warnings": problems,
    }
    if weekly:
        data["weekly"], wprob = weekly_json(*weekly)
        # how MIRA's weekly figures compare with Finance's: this year so far, and the whole of last year
        ends = data["weekly"]["end"]
        last_year = str(int(ends[-1][:4]) - 1)
        checks = [vs_finance(weekly[0], int(ends[-1][:4]), ends[-1])]
        ly = [e for e in ends if e.startswith(last_year)]
        if ly:
            checks.append(vs_finance(weekly[0], int(last_year), f"{last_year}-12-31"))
        data["weekly"]["vs_finance"] = [c for c in checks if c]
        problems += wprob
        if wprob:
            print("Warning, weekly figures:\n  " + "\n  ".join(wprob[:20]))
        print(f"Weekly figures: {len(data['weekly']['end'])} weeks, to {data['weekly']['end'][-1]}.")
    OUT.write_text(json.dumps(data, separators=(",", ":")), encoding="utf-8")
    yr = months[-1][:4]
    ytd = sum(v for m, v in zip(months, total_mvr) if m.startswith(yr))
    print(f"OK. {months[0]} to {months[-1]} ({len(months)} months), dollar figures from {usd_from}. "
          f"{yr} so far MVR {ytd / 1e9:,.2f} bn. Wrote {OUT.name}.")


if __name__ == "__main__":
    main()
