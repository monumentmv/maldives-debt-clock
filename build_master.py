"""
Builds one clean master dataset from every source the site uses:
  MMA Statistics Database (data.json), MIRA revenue files (inputs/mira/) and
  the Ministry of Finance Weekly Fiscal Developments (data/wfd/) and
  Maldives Police Service case counts (data/police/).

Runs every morning in GitHub Actions, after the other scripts. Writes:
  data/master/series_catalogue.csv     one row per series: clean name, unit, source, dates covered
  data/master/all_series_long.csv      every observation: series, date, value
  data/master/weekly_budget_tables.csv every row of every weekly budget table, cleaned
  dist/maldives_public_finance.xlsx    the same, in one Excel workbook with a sheet per frequency
  dist/maldives_public_finance_csv.zip the CSV files in one download
When the data has changed, the dist/ files are also copied to downloads/, which every page footer links to.

To run it yourself:
    pip install openpyxl
    python build_master.py
"""
import csv
import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_revenue as BR

HERE = Path(__file__).resolve().parent
OUT = HERE / "data" / "master"
DIST = HERE / "dist"
SITE = "https://mvdebtclock.org"
RELEASE = SITE + "/downloads/"
DOWNLOADS = HERE / "downloads"

# ---------------------------------------------------------------- MMA series, with clean names
# key in data.json: (series id, name, category, unit, how to convert)
MMA = {
    "total": ("mma.debt.total", "Public and publicly guaranteed debt, total", "Public debt", "MVR million", "m"),
    "domestic": ("mma.debt.domestic", "Public debt, domestic, total", "Public debt", "MVR million", "m"),
    "domestic_cg": ("mma.debt.domestic_central_government", "Public debt, domestic, central government", "Public debt", "MVR million", "m"),
    "domestic_guaranteed": ("mma.debt.domestic_guaranteed", "Public debt, domestic, government-guaranteed", "Public debt", "MVR million", "m"),
    "external": ("mma.debt.external", "Public debt, external, total", "Public debt", "MVR million", "m"),
    "external_cg": ("mma.debt.external_central_government", "Public debt, external, central government", "Public debt", "MVR million", "m"),
    "external_guaranteed": ("mma.debt.external_guaranteed", "Public debt, external, government-guaranteed", "Public debt", "MVR million", "m"),
    "debt_to_gdp": ("mma.debt.total_pct_gdp", "Public debt, total, as a share of GDP", "Public debt", "% of GDP", "pct"),
    "domestic_to_gdp": ("mma.debt.domestic_pct_gdp", "Public debt, domestic, as a share of GDP", "Public debt", "% of GDP", "pct"),
    "external_to_gdp": ("mma.debt.external_pct_gdp", "Public debt, external, as a share of GDP", "Public debt", "% of GDP", "pct"),
    "ext_interest_q": ("mma.debt.external_interest_paid", "Interest paid on central government external debt", "Public debt", "US$ million", "m"),
    "revenue": ("mma.gov.revenue_and_grants", "Government revenue and grants", "Government finance (monthly)", "MVR million", "m"),
    "tax_revenue": ("mma.gov.tax_revenue", "Government tax revenue", "Government finance (monthly)", "MVR million", "m"),
    "expenditure": ("mma.gov.spending_total", "Government spending, recurrent and capital", "Government finance (monthly)", "MVR million", "m"),
    "recurrent_exp": ("mma.gov.spending_recurrent", "Government spending, recurrent", "Government finance (monthly)", "MVR million", "m"),
    "capital_exp": ("mma.gov.spending_capital", "Government spending, capital", "Government finance (monthly)", "MVR million", "m"),
    "salaries_pensions": ("mma.gov.salaries_wages_pensions", "Salaries, wages and pensions", "Government finance (monthly)", "MVR million", "m"),
    "interest_costs": ("mma.gov.interest_costs", "Interest and other financing costs", "Government finance (monthly)", "MVR million", "m"),
    "subsidies": ("mma.gov.grants_subsidies", "Grants, contributions and subsidies", "Government finance (monthly)", "MVR million", "m"),
    "gdp": ("mma.economy.gdp_nominal", "Gross domestic product at market prices, nominal", "Economy", "MVR million", "m"),
    "population": ("mma.economy.population", "Population, including foreign residents", "Economy", "people", "x"),
    "cpi": ("mma.prices.cpi_national", "Consumer price index, national", "Prices and exchange rate", "index", "x"),
    "usd_rate": ("mma.prices.usd_rate", "Exchange rate, rufiyaa per US dollar (MMA reference rate)", "Prices and exchange rate", "MVR per US$", "x"),
    "imports_goods": ("mma.trade.imports_goods", "Imports of goods, total", "Trade and fuel", "US$ million", "m"),
    "fuel_imports": ("mma.trade.fuel_imports", "Imports of petroleum products", "Trade and fuel", "US$ million", "m"),
    "fuel_petrol": ("mma.trade.fuel_petrol", "Imports of petrol", "Trade and fuel", "US$ million", "m"),
    "fuel_diesel": ("mma.trade.fuel_diesel", "Imports of diesel (marine gas oil)", "Trade and fuel", "US$ million", "m"),
    "fuel_other": ("mma.trade.fuel_other", "Imports of other petroleum products", "Trade and fuel", "US$ million", "m"),
    "crude_price": ("mma.trade.crude_oil_price", "Crude oil price, average of Brent, Dubai and WTI (World Bank)", "Trade and fuel", "US$ per barrel", "x"),
}
MMA_NOTES = {
    "gdp": "The latest year is an MMA projection.", "population": "Includes foreign residents. The latest year is a projection.",
    "debt_to_gdp": "Uses MMA's GDP figures, projected for the latest year.", "usd_rate": "End of month.",
    "fuel_imports": "From April 2015 covers only fuel sold within Maldives.",
}

# ---------------------------------------------------------------- weekly budget measures
WFD = [  # column in wfd_headline.csv: clean name
    ("revenue_and_grants", "Revenue and grants"), ("tax_revenue", "Tax revenue"), ("non_tax_revenue", "Non-tax revenue"),
    ("grants", "Grants received"), ("tourism_gst", "Tourism GST"), ("general_gst", "GST outside tourism"),
    ("import_duties", "Import duties"), ("green_tax", "Green tax"), ("rent_from_resorts", "Rent from resorts"),
    ("soe_dividends", "Dividends from state companies"), ("total_expenditure", "Spending, recurrent and capital"),
    ("recurrent_expenditure", "Spending, recurrent"), ("capital_expenditure", "Spending, capital"),
    ("salaries_wages_pensions", "Salaries, wages and pensions"), ("subsidies", "Subsidies"), ("aasandha", "Aasandha health insurance"),
    ("interest_costs", "Interest and other financing costs"), ("primary_balance", "Primary balance"),
    ("overall_balance", "Overall balance (revenue minus spending)"), ("loan_repayment", "Loan repayments"),
    ("transfers_to_sdf", "Transfers to the Sovereign Development Fund"), ("psip", "Public sector investment programme (PSIP)"),
]
VARIANTS = [("", "ytd", "year to date"), ("_approved", "budget", "full-year approved budget"), ("_last_year", "same_date_last_year", "same date a year earlier")]
TABLES = {"summary": "Summary of government finances", "revenue": "Revenue", "expenditure": "Spending",
          "psip": "Public sector investment by function", "agencies": "Budget by government office", "securities": "Government securities outstanding"}

ACRONYMS = {"GST", "TGST", "BPT", "PSIP", "SOE", "SOES", "USD", "MVR", "CSR", "SDF", "WHT", "NWT", "RDC"}
PROPER = ["Maldives", "Malé", "Sovereign Development Fund", "Public Sector Investment Program", "Aasandha", "Judiciary Sectoral Grant",
          "Indira Gandhi Memorial Hospital", "Hulhumalé", "Hulhumale", "Brent", "Dubai", "WTI", "US dollar"]
NAME_FIX = {"gst (tourism sector)": "Tourism GST", "gst (non-tourism sector)": "GST outside tourism", "others": "Other revenue",
            "total": "Total", "fines": "Fines and penalties"}


def sentence(s):
    s = re.sub(r"\s+", " ", re.sub(r"\d+/", "", str(s))).strip(" :")
    words = s.split(" ")
    out = []
    for i, w in enumerate(words):
        bare = re.sub(r"[^A-Za-z]", "", w)
        if bare.upper() in ACRONYMS and (bare.isupper() or bare.upper() in {"GST", "TGST", "PSIP", "SOE"}):
            out.append(w.upper() if bare.isupper() else re.sub(bare, bare.upper(), w))
        else:
            out.append(w.lower())
    s = " ".join(out)
    for proper in PROPER:
        s = re.sub(re.escape(proper), proper, s, flags=re.I)
    return s[:1].upper() + s[1:]


def slug(s):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")[:60]


def month_end(ym):
    y, m = map(int, ym.split("-"))
    return (date(y + (m == 12), m % 12 + 1, 1).toordinal() - 1)


def iso_month_end(ym):
    return date.fromordinal(month_end(ym)).isoformat()


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- collect
def collect():
    series, obs = {}, []   # series_id -> metadata; obs: (series_id, date, value)

    def add_series(sid, name, category, source, ref, unit, freq, notes=""):
        series[sid] = {"order": len(series), "series_id": sid, "name": name, "category": category, "source": source, "source_reference": ref,
                       "unit": unit, "frequency": freq, "notes": notes}

    # MMA
    data = json.loads((HERE / "data.json").read_text(encoding="utf-8"))
    for key in MMA:
        s = data.get("series", {}).get(key)
        if not s:
            continue
        sid, name, cat, unit, conv = MMA[key]
        add_series(sid, name, cat, "Maldives Monetary Authority (MMA) Statistics Database", f"MMA series {s.get('id')}",
                   unit, (s.get("frequency") or "").lower(), MMA_NOTES.get(key, ""))
        for p in s.get("points", []):
            v = num(p.get("value"))
            if v is None:
                continue
            v = v / 1e6 if conv == "m" else v * 100 if conv == "pct" else v
            obs.append((sid, p["date"][:10], round(v, 4)))

    # MIRA
    sheets = {}
    for f in sorted((HERE / "inputs" / "mira").glob("*.xlsx")):
        s = BR.read_sheet(f)
        if s:
            sheets[s["kind"]] = s
    names = {}
    for kind in ("all", "usd"):
        for label in (sheets.get(kind) or {}).get("items", {}):
            names.setdefault(BR.norm(label), label)
    for kind, prefix, unit, what in (("all", "mira.collected", "MVR million", "collected, all currencies"),
                                     ("usd", "mira.paid_in_usd", "US$ million", "paid in US dollars")):
        S = sheets.get(kind)
        if not S:
            continue
        last = max((m for m, v in S["total"].items() if v), default="")
        lines = [("TOTAL", S["total"])] + list(S["items"].items())
        for label, vals in lines:
            n = BR.norm(label)
            clean = "All MIRA revenue" if n == "total" else NAME_FIX.get(n) or sentence(names.get(n, label))
            sid = f"{prefix}.{'total' if n == 'total' else slug(clean)}"
            group = dict(BR.GROUPS).get(BR.GROUP_OF.get(n, "other"), "Other taxes and fees") if n != "total" else "Total"
            add_series(sid, f"{clean}, {what}", "Revenue collected by MIRA (monthly)", "Maldives Inland Revenue Authority (MIRA)",
                       f"MIRA revenue series, line '{label.strip()}'", unit, "monthly", f"Group on the site: {group}.")
            for ym, v in vals.items():
                if v is None or ym > last or (kind == "usd" and not v and ym < "2011-01"):
                    continue
                obs.append((sid, iso_month_end(ym), round(v / 1e6, 4)))

    # Ministry of Finance, weekly
    head = list(csv.DictReader(open(HERE / "data" / "wfd" / "wfd_headline.csv", encoding="utf-8"))) if (HERE / "data" / "wfd" / "wfd_headline.csv").exists() else []
    for col, name in WFD:
        for suffix, vkey, vname in VARIANTS:
            sid = f"mof.weekly.{col}.{vkey}"
            add_series(sid, f"{name}, {vname}", "Weekly budget (Ministry of Finance)", "Ministry of Finance, Weekly Fiscal Developments",
                       f"Weekly Fiscal Developments, table 1 or 2 or 3, '{name}'", "MVR million", "weekly",
                       "Running total from 1 January to the report date." if vkey == "ytd" else
                       "As printed in each report." if vkey == "budget" else "As printed in each report, for the same date a year earlier.")
            for h in head:
                v = num(h.get(col + suffix))
                if v is not None:
                    obs.append((sid, h["as_at"], round(v, 4)))
    for col, name in (("securities_total", "Government securities outstanding, total"), ("securities_domestic", "Government securities outstanding, domestic"),
                      ("securities_external", "Government securities outstanding, external")):
        sid = f"mof.weekly.{col}"
        add_series(sid, name, "Weekly budget (Ministry of Finance)", "Ministry of Finance, Weekly Fiscal Developments",
                   "Weekly Fiscal Developments, government securities table", "MVR million", "fortnightly",
                   "Dated by the report it appears in. The table itself is as of a few days earlier.")
        for h in head:
            v = num(h.get(col))
            if v is not None:
                obs.append((sid, h["as_at"], round(v, 4)))

    # Maldives Police Service, monthly case counts
    pol = HERE / "data" / "police" / "crime_monthly.csv"
    if pol.exists():
        prow = [r for r in csv.DictReader(open(pol, encoding="utf-8")) if r["month"] != "0"]
        totals = defaultdict(int)
        for r in prow:
            totals[(int(r["year"]), int(r["month"]))] += int(r["cases"])
        add_series("mps.cases.total", "Cases reported to the police, all categories", "Cases reported to the police (monthly)",
                   "Maldives Police Service, crime statistics", "police.gov.mv/crime-statistics, total of all categories", "cases", "monthly",
                   "Reported cases, not convictions. Includes traffic accidents and lost items.")
        for (y, m), n in sorted(totals.items()):
            obs.append(("mps.cases.total", iso_month_end(f"{y}-{m:02d}"), n))
        for c in sorted({r["category"] for r in prow}):
            sid = f"mps.cases.{slug(c)}"
            add_series(sid, f"Cases reported to the police, {c.lower()}", "Cases reported to the police (monthly)",
                       "Maldives Police Service, crime statistics", f"police.gov.mv/crime-statistics, category '{c}'", "cases", "monthly",
                       "Reported cases, not convictions.")
            for r in prow:
                if r["category"] == c:
                    obs.append((sid, iso_month_end(f"{int(r['year'])}-{int(r['month']):02d}"), int(r["cases"])))

    # drop series with no data, add coverage
    by = defaultdict(list)
    for sid, d, v in obs:
        by[sid].append((d, v))
    series = {k: v for k, v in series.items() if by.get(k)}
    for sid, m in series.items():
        ds = sorted(d for d, _ in by[sid])
        m.update(first_date=ds[0], last_date=ds[-1], observations=len(ds))
    obs.sort(key=lambda o: (o[0], o[1]))
    return series, obs


def weekly_tables():
    path = HERE / "data" / "wfd" / "wfd_long.csv"
    if not path.exists():
        return []
    out = []
    for r in csv.DictReader(open(path, encoding="utf-8")):
        label = r["label"]
        if r["table"] == "agencies":
            label = re.sub(r"^\d+\s+", "", label)
        if r["table"] == "psip":
            label = re.sub(r"^\d+\.\s*", "", label)
        label = re.sub(r"^[A-G]\s+(?=[A-Z])", "", label.strip())          # row letters A to G in table 1
        label = re.sub(r"\s*\([A-G][+-][A-G]\)", "", label)              # formulas like (C+D)
        out.append({"report_date": r["as_at"], "table": TABLES.get(r["table"], r["table"]), "line": sentence(label) if label.isupper() or r["table"] != "agencies" else label,
                    "approved_budget": r["approved"], "same_date_last_year": r["last_year"], "year_to_date": r["this_year"],
                    "listed_separately": "yes" if r.get("memo") == "1" else "", "all_values_as_printed": r["all_values"]})
    return out


# ---------------------------------------------------------------- write
CAT_COLS = ["series_id", "name", "category", "unit", "frequency", "first_date", "last_date", "observations", "source", "source_reference", "notes"]


def write_csvs(series, obs, tables):
    OUT.mkdir(parents=True, exist_ok=True)
    files = {}
    buf = io.StringIO(); w = csv.DictWriter(buf, fieldnames=CAT_COLS, lineterminator="\n"); w.writeheader()
    for sid in sorted(series, key=lambda s: series[s]["order"]):
        w.writerow({k: series[sid].get(k, "") for k in CAT_COLS})
    files["series_catalogue.csv"] = buf.getvalue()
    buf = io.StringIO(); w = csv.writer(buf, lineterminator="\n"); w.writerow(["series_id", "name", "date", "value", "unit"])
    for sid, d, v in obs:
        w.writerow([sid, series[sid]["name"], d, f"{v:g}" if abs(v) < 1e15 else v, series[sid]["unit"]])
    files["all_series_long.csv"] = buf.getvalue()
    buf = io.StringIO()
    cols = ["report_date", "table", "line", "approved_budget", "same_date_last_year", "year_to_date", "listed_separately", "all_values_as_printed"]
    w = csv.DictWriter(buf, fieldnames=cols, lineterminator="\n"); w.writeheader(); w.writerows(tables)
    files["weekly_budget_tables.csv"] = buf.getvalue()
    changed = False
    for name, text in files.items():
        p = OUT / name
        if not p.exists() or p.read_text(encoding="utf-8") != text:
            p.write_text(text, encoding="utf-8")
            changed = True
    return files, changed


HEAD_FILL = PatternFill("solid", fgColor="072F40")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def sheet_table(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    for i, wdt in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt


def wide(ws, sids, series, obs_by):
    dates = sorted({d for s in sids for d in obs_by[s]})
    header = ["Date"] + [f"{series[s]['name']} ({series[s]['unit']})" for s in sids]
    sheet_table(ws, header, [[d] + [obs_by[s].get(d) for s in sids] for d in dates], [12] + [22] * len(sids))
    ws.row_dimensions[1].height = 75
    ws.append([]); ws.append(["Series IDs"] + sids)


def write_xlsx(series, obs, tables, built):
    DIST.mkdir(exist_ok=True)
    obs_by = defaultdict(dict)
    for sid, d, v in obs:
        obs_by[sid][d] = v
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Read me"
    first = lambda c: min((series[s]["first_date"] for s in series if series[s]["category"].startswith(c)), default="")
    lines = [
        ("Maldives public finance data", True),
        (f"Compiled by the Maldives National Debt Clock ({SITE}) from official publications. Built {built} UTC.", False),
        ("", False),
        ("What is in this workbook", True),
        ("Catalogue: every series, with its clean name, unit, frequency, dates covered and original source.", False),
        ("Public debt (quarterly), Monthly, Annual and Weekly budget: the same data laid out with one column per series and one row per date.", False),
        ("Weekly budget tables: every line of every table in the Ministry of Finance's weekly reports.", False),
        ("All series (long): every observation in one list, the easiest form to load into other software.", False),
        ("", False),
        ("Sources", True),
        ("Maldives Monetary Authority (MMA) Statistics Database, https://database.mma.gov.mv", False),
        ("Maldives Inland Revenue Authority (MIRA), revenue series, https://www.mira.gov.mv/Publications/Categories/310", False),
        ("Ministry of Finance, Weekly Fiscal Developments, https://www.finance.gov.mv/publications/statistical-releases/weekly-fiscal-developments", False),
        ("Maldives Police Service, crime statistics, https://www.police.gov.mv/crime-statistics", False),
        ("", False),
        ("Units and conventions", True),
        ("Money is in millions of rufiyaa (MVR million) or millions of US dollars (US$ million), as shown for each series. Ratios to GDP are in percent.", False),
        ("Dates are the last day of the month, quarter or year for MMA and MIRA series, and the report date for weekly series.", False),
        ("Weekly budget figures are running totals from 1 January and are provisional. MIRA's figures are collections, recorded when money comes in.", False),
        ("Figures are as published by each source and can be revised. They are not adjusted for inflation.", False),
        ("", False),
        ("Using this data", True),
        ("Please credit the original publishers (MMA, MIRA, the Ministry of Finance and the Maldives Police Service) and the Maldives National Debt Clock. "
         "Errors can be reported through the contact form on the site's About page.", False),
        (f"Latest version: {RELEASE}maldives_public_finance.xlsx", False),
    ]
    for text, bold in lines:
        ws.append([text]); ws.cell(ws.max_row, 1).font = Font(bold=bold, size=14 if bold and ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 140

    ws = wb.create_sheet("Catalogue")
    sheet_table(ws, [c.replace("_", " ").capitalize() for c in CAT_COLS],
                [[series[s].get(c, "") for c in CAT_COLS] for s in sorted(series, key=lambda s: series[s]["order"])],
                [38, 55, 32, 14, 11, 12, 12, 12, 40, 40, 50])
    groups = [("Public debt (quarterly)", lambda m: m["frequency"] == "quarterly"),
              ("Monthly", lambda m: m["frequency"] == "monthly"),
              ("Annual", lambda m: m["frequency"] == "annual"),
              ("Weekly budget", lambda m: m["frequency"] in ("weekly", "fortnightly"))]
    for title, test in groups:
        sids = [s for s in sorted(series, key=lambda s: series[s]["order"]) if test(series[s])]
        if sids:
            wide(wb.create_sheet(title), sids, series, obs_by)
    if tables:
        ws = wb.create_sheet("Weekly budget tables")
        cols = ["report_date", "table", "line", "approved_budget", "same_date_last_year", "year_to_date", "listed_separately"]
        sheet_table(ws, [c.replace("_", " ").capitalize() for c in cols], [[t["report_date"], t["table"], t["line"], num(t["approved_budget"]),
                    num(t["same_date_last_year"]), num(t["year_to_date"]), t["listed_separately"]] for t in tables], [12, 34, 50, 16, 18, 14, 16])
    ws = wb.create_sheet("All series (long)")
    sheet_table(ws, ["Series id", "Name", "Date", "Value", "Unit"], [[s, series[s]["name"], d, v, series[s]["unit"]] for s, d, v in obs], [38, 60, 12, 16, 14])
    wb.properties.creator = "Maldives National Debt Clock"
    wb.save(DIST / "maldives_public_finance.xlsx")


def main():
    built = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    series, obs = collect()
    tables = weekly_tables()
    files, changed = write_csvs(series, obs, tables)
    write_xlsx(series, obs, tables, built)
    with zipfile.ZipFile(DIST / "maldives_public_finance_csv.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in files.items():
            z.writestr(name, text)
        z.writestr("README.txt", "Maldives public finance data, compiled by the Maldives National Debt Clock from MMA, MIRA, Ministry of Finance and Maldives Police Service publications.\n"
                   f"Built {built} UTC. See series_catalogue.csv for names, units and sources. Money is in millions of MVR or US$ as stated.\n")
    meta = {"built_at": built, "series": len(series), "observations": len(obs), "weekly_table_rows": len(tables),
            "first_date": min(m["first_date"] for m in series.values()), "last_date": max(m["last_date"] for m in series.values()),
            "xlsx": RELEASE + "maldives_public_finance.xlsx", "zip": RELEASE + "maldives_public_finance_csv.zip"}
    if changed or not (DOWNLOADS / "maldives_public_finance.xlsx").exists():
        DOWNLOADS.mkdir(exist_ok=True)
        for f in ("maldives_public_finance.xlsx", "maldives_public_finance_csv.zip"):
            (DOWNLOADS / f).write_bytes((DIST / f).read_bytes())
    (OUT / "about.json").write_text(json.dumps(meta, indent=1), encoding="utf-8") if changed or not (OUT / "about.json").exists() else None
    print(f"Master data: {len(series)} series, {len(obs):,} observations, {len(tables):,} weekly table rows. "
          f"{'Changed' if changed else 'No change'} since the last build.")


if __name__ == "__main__":
    main()
